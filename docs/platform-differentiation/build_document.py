"""Build the editable report and PDF locally; no external services required."""
from pathlib import Path
import html
import re
import subprocess
import tempfile


def write_word(source_html, destination):
    """Requires python-docx and beautifulsoup4; keep them outside the app env."""
    from bs4 import BeautifulSoup, NavigableString
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(16)
    section.left_margin = section.right_margin = Mm(17)
    section.header_distance = section.footer_distance = Mm(8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet'):
        style = doc.styles[name]
        style.font.name = 'Noto Sans CJK KR'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
        style.font.size = Pt(9)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 23), ('Heading 1', 15), ('Heading 2', 11)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('393665')
    header = section.header.paragraphs[0]
    header.add_run('MeDIAuto Studio  |  플랫폼 차별성 분석').font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('2026-09-14 · 문서 v1.0    |    ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    def rich(p, node, bold=False):
        for child in node.children:
            if isinstance(child, NavigableString):
                r = p.add_run(str(child))
                r.bold = bold
            elif child.name == 'a':
                link = OxmlElement('w:hyperlink')
                link.set(qn('r:id'), p.part.relate_to(child['href'], RT.HYPERLINK, is_external=True))
                run = OxmlElement('w:r')
                props = OxmlElement('w:rPr')
                color = OxmlElement('w:color'); color.set(qn('w:val'), '266A88'); props.append(color)
                underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
                run.append(props)
                txt = OxmlElement('w:t'); txt.text = child.get_text(); run.append(txt)
                link.append(run); p._p.append(link)
            else:
                rich(p, child, bold or child.name == 'strong')

    soup = BeautifulSoup(source_html, 'html.parser')
    for node in soup.body.children:
        if isinstance(node, NavigableString):
            continue
        if node.name in ('h1', 'h2', 'h3'):
            p = doc.add_paragraph(style={'h1': 'Title', 'h2': 'Heading 1', 'h3': 'Heading 2'}[node.name])
            rich(p, node)
            if 'page-break-before' in node.get('style', ''):
                p.paragraph_format.page_break_before = True
        elif node.name == 'p':
            rich(doc.add_paragraph(), node)
        elif node.name == 'ul':
            for li in node.find_all('li', recursive=False):
                rich(doc.add_paragraph(style='List Bullet'), li)
        elif node.name == 'table':
            rows = node.find_all('tr')
            count = len(rows[0].find_all(['td', 'th']))
            table = doc.add_table(rows=0, cols=count)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [38, 138] if count == 2 else [32, 72, 72]
            for col, width in zip(table.columns, widths):
                col.width = Mm(width)
            for index, source_row in enumerate(rows):
                row = table.add_row()
                trpr = row._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if index == 0:
                    trpr.append(OxmlElement('w:tblHeader'))
                for cell, source_cell, width in zip(row.cells, source_row.find_all(['td', 'th']), widths):
                    cell.width = Mm(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.line_spacing = 1.05
                    rich(p, source_cell, index == 0)
                    for run in p.runs:
                        run.font.size = Pt(8)
                    if index == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8E7F2')
                        cell._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
    doc.core_properties.title = 'MeDIAuto Studio 플랫폼 차별성 분석'
    doc.core_properties.subject = '특허 기술 후보와 구현 근거 및 경쟁 플랫폼 비교'
    doc.core_properties.author = 'MeDIAuto Studio'
    doc.save(destination)

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'MeDIAuto_differentiation_2026-09-14.md'


def inline(value):
    value = html.escape(value)
    value = re.sub(r'\[([^\]]+)\]\((https://[^)]+)\)', r'<a href="\2">\1</a>', value)
    value = re.sub(r'`([^`]+)`', r'<code>\1</code>', value)
    return re.sub(r'\*\*([^*]+)\*\*', r'<strong>\1</strong>', value)


def render():
    lines = SOURCE.read_text().splitlines()
    blocks = []
    index = 0
    sections = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                cells = [c.strip() for c in lines[index].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c) for c in cells):
                    rows.append(cells)
                index += 1
            parts = ['<table width="100%" cellspacing="0" cellpadding="6">']
            for row_index, cells in enumerate(rows):
                tag = 'th' if row_index == 0 else 'td'
                if row_index == 0:
                    parts.append('<thead>')
                elif row_index == 1:
                    parts.append('<tbody>')
                parts.append('<tr>' + ''.join(f'<{tag}>{inline(c)}</{tag}>' for c in cells) + '</tr>')
                if row_index == 0:
                    parts.append('</thead>')
            parts.append('</tbody></table>')
            blocks.append(''.join(parts))
            continue
        if line.startswith('- '):
            parts = ['<ul>']
            while index < len(lines) and lines[index].startswith('- '):
                parts.append('<li>' + inline(lines[index][2:]) + '</li>')
                index += 1
            blocks.append(''.join(parts) + '</ul>')
            continue
        match = re.match(r'^(#{1,3}) (.*)', line)
        if match:
            level = len(match[1])
            attr = ''
            if level == 2:
                sections += 1
                if sections > 1:
                    attr = ' style="page-break-before: always"'
            blocks.append(f'<h{level}{attr}>{inline(match[2])}</h{level}>')
        else:
            blocks.append('<p>' + inline(line) + '</p>')
        index += 1
    css = '''
    @page {size:A4; margin:17mm 17mm 16mm 17mm;}
    body {font-family:"Noto Sans CJK KR",sans-serif;font-size:9.5pt;color:#243247;line-height:1.38;}
    h1 {font-size:23pt;color:#393665;margin:0 0 12pt;}
    h2 {font-size:15pt;color:#393665;margin:12pt 0 10pt;page-break-after:avoid;}
    h3 {font-size:11.5pt;color:#216872;margin:14pt 0 6pt;page-break-after:avoid;}
    p {margin:0 0 8pt;orphans:2;widows:2;}
    table {border-collapse:collapse;table-layout:fixed;margin:6pt 0 10pt;font-size:8.3pt;width:100%;}
    th {background:#e8e7f2;color:#302c59;font-weight:bold;}
    td,th {border:0.5pt solid #cbd2dc;padding:6pt;text-align:left;vertical-align:top;word-break:break-all;}
    tr {page-break-inside:avoid;} thead {display:table-header-group;}
    a {color:#266a88;text-decoration:underline;}
    code {font-family:"Noto Sans CJK KR",sans-serif;font-size:8pt;word-break:break-all;}
    ul {padding-left:15pt;margin:5pt 0 10pt;} li {margin-bottom:5pt;}
    '''
    return '<!DOCTYPE html><html lang="ko"><head><meta charset="utf-8"><title>MeDIAuto Studio 플랫폼 차별성 분석</title><style>' + css + '</style></head><body>' + '\n'.join(blocks) + '</body></html>'


if __name__ == '__main__':
    source_html = SOURCE.with_suffix('.html')
    source_html.write_text(render(), encoding='utf-8')
    write_word(source_html.read_text(), SOURCE.with_suffix('.docx'))
    with tempfile.TemporaryDirectory(prefix='mediauto-report-lo-') as profile:
        base = ['libreoffice', '-env:UserInstallation=' + Path(profile).as_uri(), '--headless']
        subprocess.run(base + ['--convert-to', 'pdf:writer_pdf_Export', '--outdir', str(ROOT), str(SOURCE.with_suffix('.docx'))], check=True)
    print('Created HTML, DOCX and PDF:', SOURCE.stem)
