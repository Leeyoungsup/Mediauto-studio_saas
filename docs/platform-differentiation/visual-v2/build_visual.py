"""Build the visual brief. Requires python-docx, Pillow and LibreOffice.

Example: PYTHONPATH=/tmp/mediauto-report-deps python build_visual.py
Existing screenshots are embedded unchanged; no AI-generated product UI.
"""
from pathlib import Path
import html
import json
import subprocess
import tempfile
from PIL import Image
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
STEM = 'MeDIAuto_visual_differentiation_v2'
NAVY, PURPLE, TEAL = '302D59', '6959D9', '197C82'

SOURCES = [
 ('H1', 'HALO AI', 'https://indicalab.com/halo-ai/'),
 ('H2', 'HALO Link', 'https://indicalab.com/halo-link/'),
 ('A1', 'Aiforia Create', 'https://www.aiforia.com/aiforia-create'),
 ('V1', 'Visiopharm · Phenotyping', 'https://visiopharm.com/resources/advanced-phenotyping-of-the-tumor-immune-microenvironment/'),
 ('V2', 'Visiopharm · AI Author 활용', 'https://visiopharm.com/blog/alimentivs-digital-pathology-innovation-with-visiopharm-an-interview-with-dr-pavine-lefevre/'),
 ('P1', 'Proscia · Concentriq', 'https://proscia.com/platform/concentriq-digital-pathology-platform'),
 ('P2', 'Concentriq LS 브로슈어', 'https://info.proscia.com/hubfs/Life-sciences/Proscia_ConcentriqLS_Brochure.pdf'),
 ('R1', 'PathPresenter', 'https://www.pathpresenter.com/about-pathpresenter/'),
 ('Q1', 'QuPath', 'https://qupath.github.io/'),
 ('Q2', 'QuPath · Scripting', 'https://qupath.github.io/workshop-intro/chapters/extras/how_do_i_write_scripts.html'),
]

PAGES = [
 dict(kicker='VISUAL BRIEF  /  2026.09.14  /  v2.0', title='분석 결과를, 사람이 검토하는 작업으로',
      subtitle='MeDIAuto Studio  |  특허 기술 후보와 실제 화면으로 보는 플랫폼 차별성',
      image='screen-multiview.png', kind='screen',
      notes=[('같이 보기','다중 화면에서 서로 다른 슬라이드와 분석 결과를 함께 확인합니다.'),('위치에서 검토','원본 위의 세포 표시와 가상염색 비교를 통해 결과의 위치를 살펴봅니다.'),('작업으로 연결','조직·세포 주석과 패치 검수 기능을 같은 플랫폼에서 제공합니다.')],
      takeaway='핵심 제안  ·  정량 AI → 가상염색 비교 → 주석·검수로 이어지는 구체적인 작업 경험',
      caption='실제 화면 S1 · 저장소의 기존 비식별 스크린샷. 화면 명칭·모델 표시는 촬영 당시 기준이며 현재 UI와 다를 수 있습니다.'),
 dict(kicker='01  /  PLATFORM WORKFLOW', title='차별성은 기능 사이의 연결에서 드러납니다',
      subtitle='원본을 열고, 결과를 살피고, 필요한 작업을 사람이 검수하는 흐름',
      image='infographic-workflow.png',kind='infographic',
      notes=[('공통 기반','슬라이드 좌표를 기준으로 원본·결과·주석의 위치를 연결합니다.'),('선택적 작업','분석·가상염색·주석은 목적에 따라 선택합니다. 필수 순서를 뜻하지 않습니다.'),('사람의 검수','수정·검수 데이터를 보관합니다. 자동 재학습까지 완료된 구조는 아닙니다.')],
      takeaway='제품 포지션  ·  기관별 병리 AI 검토·라벨링 업무를 연결하는 웹 기반 작업 플랫폼',
      caption='생성 인포그래픽 G1 · 코드와 특허 초안을 설명하는 개념도. 실제 제품 화면이나 성능 측정 자료가 아닙니다.'),
 dict(kicker='02  /  VIRTUAL STAIN  /  특허 기술 후보 ③', title='가상염색 결과를 원본 위치에서 비교합니다',
      subtitle='생성 결과를 별도 파일로만 보관하지 않고, 정렬된 비교 화면으로 전달',
      image='screen-virtual-stain.png',kind='screen',
      notes=[('중앙 분할선','원본과 생성 영상을 나누는 경계를 이동하며 같은 위치를 비교합니다.'),('오른쪽 제어','Overlay와 Split View를 켜고 끄는 조작을 확인할 수 있습니다.'),('처리 구조','패치 추론·겹침 혼합·타일 피라미드·정렬 표시가 연결됩니다.')],
      takeaway='차별화 제안  ·  가상염색 생성과 탐색·비교를 하나의 작업 흐름으로 제공',
      caption='실제 화면 S2 · 화면 예시는 IHC → Virtual H&E. 변환 방향을 반대로 해석하지 않습니다. 실제 검사 대체·진단 성능의 근거가 아닙니다.'),
 dict(kicker='03  /  QUANTITATIVE AI  /  특허 기술 후보 ②·④', title='숫자의 근거가 된 세포를 직접 확인하고 수정합니다',
      subtitle='원본 영상 위에서 세포별 분류와 신뢰도를 확인하는 검토 화면',
      image='screen-cell-edit.png',kind='screen',
      notes=[('중앙 편집 메뉴','선택한 세포의 클래스와 신뢰도를 확인하고 클래스 변경·삭제를 수행합니다.'),('오른쪽 결과','세포 분류와 집계가 영상 검토와 함께 제공됩니다.'),('저장 방식','정량 AI 원본과 사용자별 수정 결과를 분리해 보관하는 코드가 있습니다.')],
      takeaway='차별화 제안  ·  결과 요약에서 세포 단위 근거 확인과 사람의 수정으로 연결',
      caption='실제 화면 S3 · 스크린샷의 HER2 점수·세포 수는 예시 결과이며 정확도 검증 자료가 아닙니다. 타사에도 AI 결과 수정 기능이 있습니다.'),
 dict(kicker='04  /  SPATIAL CONTEXT  /  특허 기술 후보 ②', title='세포의 개수와 공간 분포를 함께 살펴봅니다',
      subtitle='세포 단위 표시와 공간 요약을 목적에 맞게 전환하는 검토 경험',
      image='screen-heatmap.png',kind='screen',
      notes=[('왼쪽 시각화','Spatial Heatmap 화면에서 조직 내 분포를 확인합니다.'),('원본의 위치','배경 뷰어의 세포 표시와 위치 정보를 바탕으로 결과를 검토합니다.'),('표시 최적화','현재 코드는 화면 내 결과를 선별하고 확대 수준·결과 수에 따라 표현을 조절합니다.')],
      takeaway='차별화 제안  ·  많은 결과를 표시하는 것에서, 검토 대상을 찾고 좁히는 경험으로 연결',
      caption='실제 화면 S4 · 공간 격자·표현 전환의 구현 근거는 tile-viewer.js. 정지 화면만으로 처리 속도나 수백만 결과의 실시간 성능을 증명하지 않습니다.'),
 dict(kicker='05  /  PATCH WORKFLOW  /  특허 기술 후보 ④', title='슬라이드의 필요한 영역을 검수 가능한 작업으로 나눕니다',
      subtitle='필요 영역과 패치 목록, 주석·검수·종결 상태의 연결',
      image='screen-patch-review.png',kind='screen',
      notes=[('중앙 패치 격자','선택된 패치 위치를 전체 슬라이드 맥락에서 확인합니다.'),('오른쪽 목록','각 패치의 Annotation·Review·Termination 상태를 함께 표시합니다.'),('상단 진행 상태','현재 작업 진행 상황을 확인하고 다음 검토 대상을 선택합니다.')],
      takeaway='차별화 제안  ·  주석을 그리는 기능에 작업 단위와 검수 상태를 결합',
      caption='실제 화면 S5 · 표시된 완료율은 촬영 당시 작업 상태의 예시입니다. 업무 효율 향상률이나 자동 검수 정확도를 의미하지 않습니다.'),
 dict(kicker='06  /  CELL ANNOTATION  /  특허 기술 후보 ④', title='AI 보조 결과를 사람이 편집할 수 있는 주석으로 다룹니다',
      subtitle='패치 내 세포, 클래스 목록, 작업 상태를 한 화면에서 확인',
      image='screen-cell-annotation.png',kind='screen',
      notes=[('중앙 세포 박스','세포별 위치와 분류를 보며 검토합니다.'),('오른쪽 편집','클래스 관리, BBox·Point 표시와 세포 주석 목록을 제공합니다.'),('검수와 연결','저장과 재열람, 패치 상태 및 내보내기 경로가 구현되어 있습니다.')],
      takeaway='차별화 제안  ·  AI 보조 → 사람의 수정 → 패치 검수 상태로 연결되는 업무 구성',
      caption='실제 화면 S6 · 현재 확인 범위는 주석·수정·상태 관리입니다. 수정 데이터의 자동 재학습·모델 배포는 구현 완료로 설명하지 않습니다.'),
 dict(kicker='07  /  RESPONSIVE VIEWING  /  특허 기술 후보 ①', title='화면 뒤에서는 현재 요청을 먼저 처리하도록 조정합니다',
      subtitle='스크린샷에 드러나지 않는 차별화 요소: 화면 요청과 서버 생성 우선순위의 연결',
      image='infographic-tiles.png',kind='infographic',
      notes=[('캐시 재사용','생성된 타일이 있으면 이를 반환합니다.'),('부하에 맞는 생성','캐시가 없으면 우선 요청을 등록하고 실행 수·대기열에 따라 즉시 생성 가능 여부를 판단합니다.'),('작업 양보','뷰어 활동을 고려해 일부 백그라운드 작업의 실행을 조정합니다.')],
      takeaway='검증할 효과  ·  첫 유효 화면 시간 / 이동 후 목표 해상도 시간 / AI 병행 시 p95 응답시간',
      caption='생성 인포그래픽 G2 · 흐름을 단순화한 개념도. 즉시 생성 경로에서도 우선 요청을 등록합니다. 타사 대비 속도·사용자별 공정성 보장을 뜻하지 않습니다.'),
]

COMPETITORS = [
 ('HALO 계열','AI 학습·정량 분석·공동 데이터 작업','가상염색 비교와 패치 검수를 연결하는 실제 작업으로 설명','H1, H2'),
 ('Aiforia Create','클라우드 모델 개발·검증·AI 보조 주석','기관 서버 중심 운영과 구체적인 검수 절차를 제시','A1'),
 ('Visiopharm','예제 기반 AI 학습·조직/세포 표현형 분석','브라우저에서 결과 비교와 주석 작업을 연결하는 흐름을 제시','V1, V2'),
 ('Proscia Concentriq LS','이미지·연관 데이터 관리·AI 실행 연결','사용자 수정 저장과 패치 상태 관리를 구체적으로 시연','P1, P2'),
 ('PathPresenter','기업용 이미지 관리·협업·외부 시스템 연동','자체 분석 파이프라인과 세포·패치 검수의 연결을 제시','R1'),
 ('QuPath','설치형 공개 소스 분석·주석·스크립트','계정·서버 작업·패치 검수를 포함한 웹 업무 제공 범위를 설명','Q1, Q2'),
]


def shade(cell, fill):
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(el)
    margins = OxmlElement('w:tcMar')
    for side in ('left', 'right'):
        edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'180');edge.set(qn('w:type'),'dxa');margins.append(edge)
    cell._tc.get_or_add_tcPr().append(margins)


def para(parent, text='', size=11, color=NAVY, bold=False, after=5):
    p = parent.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def clean_cell(cell):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    p.add_run().font.size = Pt(1)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def table(parent, widths):
    t = parent.add_table(rows=1, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for col, cell, width in zip(t.columns, t.rows[0].cells, widths):
        col.width = cell.width = Mm(width); clean_cell(cell)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return t


def title(doc, kicker, heading, sub):
    para(doc,kicker,9,TEAL,True,5)
    para(doc,heading,23,NAVY,True,5)
    para(doc,sub,11,'576174',False,8)


def bar(doc,text):
    t = table(doc,[269]); shade(t.cell(0,0),'EFEDF9')
    para(t.cell(0,0),text,11,NAVY,True,6)


def link(p, label, url):
    el = OxmlElement('w:hyperlink'); el.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),TEAL);props.append(col)
    run.append(props);txt=OxmlElement('w:t');txt.set(qn('xml:space'),'preserve');txt.text=label;run.append(txt);el.append(run);p._p.append(el)


def build_docx():
    doc=Document();s=doc.sections[0]
    s.page_width,s.page_height=Mm(297),Mm(210)
    s.left_margin=s.right_margin=Mm(14)
    s.top_margin=Mm(12);s.bottom_margin=Mm(12)
    s.header_distance=s.footer_distance=Mm(5)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Noto Sans CJK KR'
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK KR')
        st.font.size=Pt(10);st.paragraph_format.line_spacing=1.1
        st.paragraph_format.space_after=Pt(5)
    footer=s.footer.paragraphs[0];footer.alignment=2
    footer.add_run('MeDIAuto Studio  ·  Visual Brief v2.0  |  ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for i,data in enumerate(PAGES):
        if i: doc.add_page_break()
        title(doc,data['kicker'],data['title'],data['subtitle'])
        is_graph=data['kind']=='infographic'
        t=table(doc,[211,58]);left,right=t.rows[0].cells
        p=left.add_paragraph();p.paragraph_format.space_after=Pt(0);p.alignment=1
        img=ASSETS/data['image'];w,h=Image.open(img).size
        width=min(207,126*w/h) if is_graph else 207
        p.add_run().add_picture(str(img),width=Mm(width))
        for idx,(head,body) in enumerate(data['notes'],1):
            para(right,f'{idx:02d}  {head}',11,PURPLE,True,4)
            para(right,body,10,NAVY,False,10)
        para(doc,data['caption'],8,'626B7B',False,6)
        bar(doc,data['takeaway'])
    doc.add_page_break()
    title(doc,'08  /  COMPETITIVE POSITION','경쟁사와 겹치는 기능 위에, 우리 작업 흐름을 제시합니다',
          '공식 자료 기반 비교 · 점수와 순위가 아닌 제품의 초점과 제안 방향')
    for row in range(3):
        t=table(doc,[134.5,134.5])
        for cell,data in zip(t.rows[0].cells,COMPETITORS[row*2:row*2+2]):
            vendor,focus,proposal,refs=data;shade(cell,'F3F2FA')
            para(cell,vendor,14,PURPLE,True,4)
            para(cell,'공식 확인  ·  '+focus,10,NAVY,False,4)
            para(cell,'우리 제안  ·  '+proposal,10,NAVY,True,4)
            p=para(cell,'근거  ',8,'626B7B',False,4)
            for ref in refs.split(', '):
                src=next(x for x in SOURCES if x[0]==ref);link(p,ref+'  ',src[2])
        para(doc,'',2,after=2)
    bar(doc,'차별성의 입증 단위  ·  기능 개수보다 “한 슬라이드의 검토·수정·검수 완료 과정”')
    para(doc,'타사 내부 구현이 공개 자료에 없다고 기능이 없다고 판단하지 않습니다. 연구용 제품·별도 모듈·진단용 제품의 범위는 구분합니다.',8,'626B7B')
    doc.add_page_break()
    title(doc,'09  /  PATENT → PRODUCT','네 가지 특허 기술 후보를 화면과 처리 구조에 연결했습니다',
          '색상 구분: 보라 = 코드 확인 / 회색 = 추가 구현 또는 검증 · 출원·등록 상태는 공식 서류 확인 필요')
    cards=[('① 적응형 타일 처리','화면 요청 → 부하 판단 → 타일 우선 생성','8쪽 개념도 · tiles.py / priority.py','사용자별 공정성·가중 점수식의 완전 구현'),
           ('② 정량 결과 시각화','공간 격자 → 화면 내 선별 → 표현 전환','4~5쪽 실제 화면 · tile-viewer.js','수백만 결과의 무지연 처리 실측'),
           ('③ 가상염색 비교','패치 혼합 → 피라미드 → 원본 정렬','3쪽 실제 화면 · virtual_stain.py','불확실성 지도의 검증·실제 검사와 일치도'),
           ('④ AI 보조 주석·검수','필요 영역 → 패치 → 수정·검수 상태','6~7쪽 실제 화면 · cell_annotation.py','자동 재학습·모델 검증·배포 루프')]
    for row in range(2):
        t=table(doc,[134.5,134.5])
        for cell,(heading,flow,evidence,gap) in zip(t.rows[0].cells,cards[row*2:row*2+2]):
            shade(cell,'F3F2FA');para(cell,heading,15,PURPLE,True)
            para(cell,flow,12,NAVY,True)
            para(cell,'구현 근거  '+evidence,9,TEAL)
            para(cell,'추가 확인  '+gap,9,'697185')
        para(doc,'',2,after=4)
    bar(doc,'다음 입증 과제  ·  응답시간(p95)  /  원본·생성 좌표 일치  /  검수 완료 시간  /  수정 결과 보존')
    para(doc,'특허 후보와 코드의 상세 연결·한계는 기존 8페이지 분석 문서를 참조합니다. 이 그림은 신규성·진보성 또는 타사 특허 비침해에 대한 결론이 아닙니다.',8,'626B7B')
    doc.add_page_break()
    title(doc,'10  /  EVIDENCE & ASSET NOTES','화면·개념도·비교 근거를 구분해 확인할 수 있습니다',
          '문서 v2.0 · 소스 기준 v3.1.0 / Git 4c98bf6 · 공식 자료 확인일 2026-09-14')
    t=table(doc,[134.5,134.5]);left,right=t.rows[0].cells
    para(left,'실제 화면 S1~S6',14,PURPLE,True)
    para(left,'figure/manual-redacted/의 기존 캡처를 변경 없이 사용했습니다. 사용자 정보 등은 해당 원본 폴더에서 비식별 처리한 상태입니다. 캡처 당시 화면과 현재 UI 명칭은 다를 수 있습니다.',10)
    para(left,'새로운 인포그래픽 G1~G2',14,PURPLE,True)
    para(left,'내장 이미지 생성 도구로 제작한 개념도입니다. 실제 UI·환자 영상·성능 그래프를 생성하지 않았습니다. 생성 프롬프트와 자산 출처는 같은 폴더에 보관했습니다.',10)
    para(left,'읽는 범위',14,PURPLE,True)
    para(left,'경쟁사 기능은 공식 자료의 제품·모듈 범위로 비교했습니다. 기능 연결을 차별화 제안으로 해석한 것이며 속도·정확도·독점성 우위를 증명한 것은 아닙니다.',10)
    para(left,'원본 분석 문서',14,PURPLE,True)
    para(left,'상위 폴더의 MeDIAuto_differentiation_2026-09-14.md / .docx / .pdf에 상세 구현 근거, 코드 위치 및 검증 계획이 있습니다. 특허 자료는 patent/의 출원 후보 초안을 사용했습니다.',9)
    para(right,'경쟁 플랫폼 공식 자료',14,PURPLE,True)
    for key,label,url in SOURCES:
        p=para(right,'',10,after=7);link(p,f'{key}  ·  {label}',url)
    para(doc,'파일 구성  ·  PDF: 화면 중심 열람 / Word: 본문·설명 수정 / HTML: 확대 가능한 웹 열람 / assets: 스크린샷·인포그래픽 원본',9,TEAL)
    doc.core_properties.title='MeDIAuto Studio · 플랫폼 차별성 비주얼 브리프'
    doc.core_properties.author='MeDIAuto Studio';doc.core_properties.version='2.0'
    doc.save(ROOT/(STEM+'.docx'))


def build_html():
    esc=html.escape
    parts=[]
    for data in PAGES:
        notes=''.join(f'<div><b>{i:02d} · {esc(h)}</b><p>{esc(b)}</p></div>' for i,(h,b) in enumerate(data['notes'],1))
        parts.append(f'<section><small>{esc(data["kicker"])}</small><h1>{esc(data["title"])}</h1><p class="sub">{esc(data["subtitle"])}</p><div class="visual"><a href="assets/{data["image"]}" target="_blank"><img src="assets/{data["image"]}" alt="{esc(data["title"])}"></a><aside>{notes}</aside></div><p class="caption">{esc(data["caption"])}</p><strong class="bar">{esc(data["takeaway"])}</strong></section>')
    card_parts=[]
    for vendor,focus,proposal,refs in COMPETITORS:
        links=[]
        for ref in refs.split(', '):
            url=next(s[2] for s in SOURCES if s[0]==ref)
            links.append(f'<a href="{url}">{ref}</a>')
        card_parts.append(f'<article><h2>{esc(vendor)}</h2><p>공식 확인 · {esc(focus)}</p><b>우리 제안 · {esc(proposal)}</b><p>'+ ' '.join(links)+'</p></article>')
    cards=''.join(card_parts)
    parts.append('<section><small>08 / COMPETITIVE POSITION</small><h1>경쟁사와 겹치는 기능 위에, 우리 작업 흐름을 제시합니다</h1><div class="cards">'+cards+'</div><p class="caption">공식 공개 자료 비교이며 타사 기능 부재·임상 성능 우위를 주장하지 않습니다.</p></section>')
    parts.append('<section><h1>상세 근거 및 편집 가능한 문서</h1><p><a href="'+STEM+'.pdf">전체 11페이지 PDF</a> · <a href="'+STEM+'.docx">Word 문서</a> · <a href="../MeDIAuto_differentiation_2026-09-14.html">상세 기술 분석</a></p><h2>공식 출처</h2>'+''.join(f'<p><a href="{u}">{k} · {esc(l)}</a></p>' for k,l,u in SOURCES)+'</section>')
    css='''body{margin:0;background:#eeedf5;color:#302d59;font:16px/1.55 "Noto Sans CJK KR",sans-serif}section{box-sizing:border-box;max-width:1400px;margin:24px auto;background:white;padding:44px 48px;border-radius:12px}small{color:#197c82;font-weight:bold;letter-spacing:1px}h1{font-size:32px;line-height:1.3;margin:12px 0}h2{font-size:22px;margin:0 0 10px}.sub,.caption{color:#626b7b}.visual{display:grid;grid-template-columns:4fr 1fr;gap:22px;align-items:center}.visual img{width:100%;max-height:640px;object-fit:contain}.visual aside b{color:#6959d9}.visual aside p{font-size:15px}.caption{font-size:12px}.bar{display:block;padding:14px;background:#efedf9;border-radius:7px}.cards{display:grid;grid-template-columns:1fr 1fr;gap:18px}article{background:#f3f2fa;padding:20px;border-radius:8px}a{color:#197c82}@media(max-width:850px){section{padding:24px;margin:12px}.visual,.cards{grid-template-columns:1fr}.visual aside{display:flex;gap:16px}.visual aside>div{flex:1}h1{font-size:25px}}@media print{section{page-break-after:always;margin:0;box-shadow:none}body{background:white}}'''
    (ROOT/(STEM+'.html')).write_text('<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>MeDIAuto 비주얼 브리프</title><style>'+css+'</style>'+''.join(parts)+'</html>',encoding='utf-8')
    (ROOT/'visual-content.json').write_text(json.dumps({'pages':PAGES,'competitors':COMPETITORS,'sources':SOURCES},ensure_ascii=False,indent=2),encoding='utf-8')


if __name__=='__main__':
    build_docx();build_html()
    with tempfile.TemporaryDirectory(prefix='mediauto-visual-lo-') as profile:
        subprocess.run(['libreoffice','-env:UserInstallation='+Path(profile).as_uri(),'--headless','--convert-to','pdf:writer_pdf_Export','--outdir',str(ROOT),str(ROOT/(STEM+'.docx'))],check=True)
    print(ROOT/(STEM+'.pdf'))
